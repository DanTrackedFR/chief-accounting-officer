"""Bounded binary parsing. No macro/formula execution, fetching or review authority."""
import copy
import base64
import csv
import io
import zipfile
from datetime import date,datetime
from decimal import Decimal
from orchestration.intake.sources import RawSource, Inventory, fingerprint, bounded
from local_cao.context import identity
import hashlib

MAX_BYTES=1_000_000
MAX_EXPANDED=8_000_000
MAX_CELLS=20_000
MAX_ROWS=5_000
MAX_PAGES=100
ROLES={'contract','invoice','delivery','reconciliation_current','reconciliation_comparative','movements','trial_balance','supporting_schedule','policy'}


def archive(data):
    z=zipfile.ZipFile(io.BytesIO(data));members=z.infolist()
    if len(members)>1000 or sum(m.file_size for m in members)>MAX_EXPANDED:raise ValueError('Expanded document limit')
    for m in members:
        name=m.filename.lower()
        if m.flag_bits&1 or any(x in name for x in ('vbaproject','embeddings/','activex/','scripts/')):raise ValueError('Active/embedded document unsupported')
    return z


def scalar(v):
    if isinstance(v,(date,datetime)):return v.isoformat()
    if v is None or type(v) in (str,int,float,bool):return v
    raise ValueError('Unsupported spreadsheet cell')


def _ingest(doc, company, scope):
    allowed={'id','version','format','content_base64','role','metadata','supersedes'}
    if not isinstance(doc,dict) or set(doc)-allowed or allowed-{'supersedes'}-set(doc):raise ValueError('Document envelope')
    identity(doc['id']);identity(doc['version'])
    from interfaces.public_output import public_record
    for v in (doc['id'],doc['version']):public_record({'guidance':v},route='tool_output')
    if doc['role'] not in ROLES:raise ValueError('Document role')
    meta=copy.deepcopy(doc['metadata']);bounded(meta)
    if not isinstance(meta,dict) or set(meta)-{'company_id','entity','period','currency','row_count','control_total','complete_population','framework','scope_id','period_id','calendar_id','period_role','relationship_id','controlled_export','comparator'}:raise ValueError('Document metadata')
    if (meta.get('company_id'),meta.get('entity'),meta.get('currency'))!=(company,scope['entity'],scope['currency']):raise ValueError('Document namespace/dimensions')
    if meta.get('framework') and meta['framework']!=scope['framework']:raise ValueError('Document framework')
    p=meta.get('period')
    if not isinstance(p,list) or len(p)!=2 or any(type(x) is not str or date.fromisoformat(x).isoformat()!=x for x in p) or p[0]>p[1]:raise ValueError('Document period')
    if doc['role']=='reconciliation_comparative':
        if p[1]>=scope['period_start']:raise ValueError('Comparative period overlaps')
    elif p!=[scope['period_start'],scope['reporting_period']]:raise ValueError('Document period mismatch')
    if type(doc['content_base64']) is not str or len(doc['content_base64'])>4*MAX_BYTES//3+8:raise ValueError('Document byte limit')
    data=base64.b64decode(doc['content_base64'],validate=True)
    if not data or len(data)>MAX_BYTES:raise ValueError('Document byte limit')
    fmt=doc['format'];tables=[];blocks=[];warnings=[];cells=0
    def table(name, rows):
        nonlocal cells
        if not rows:raise ValueError('Empty table')
        headers=rows[0]
        if any(type(h) is not str or not h or len(h)>120 for h in headers) or len(set(headers))!=len(headers):raise ValueError('Explicit unique headers required')
        records=[]
        for i,row in enumerate(rows[1:],2):
            if len(row)!=len(headers):raise ValueError('Incomplete table width')
            cells+=len(row)
            if cells>MAX_CELLS or i>MAX_ROWS+1:raise ValueError('Population limit; no truncation')
            records.append(dict(zip(headers,[scalar(v) for v in row])))
        tables.append(dict(name=name,rows=records))
    if fmt=='csv':
        table('table',list(csv.reader(io.StringIO(data.decode('utf-8-sig')),strict=True)))
    elif fmt=='xlsx':
        from openpyxl import load_workbook
        with archive(data):pass
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            for sheet in book:
                # Ignore potentially forged dimension caches; scan actual XML cells.
                sheet.reset_dimensions();rows=[]
                for row in sheet.iter_rows():
                    values=[]
                    for cell in row:
                        if cell.data_type in ('f','e'):raise ValueError('Formula/error cells need independently recalculated export')
                        values.append(scalar(cell.value))
                    rows.append(values)
                    if len(rows)>MAX_ROWS+1 or sum(len(x) for x in rows)>MAX_CELLS:raise ValueError('Spreadsheet limit')
                if rows:
                    width=max(len(x) for x in rows);rows=[x+[None]*(width-len(x)) for x in rows]
                    table(sheet.title,rows)
                    if sheet.sheet_state!='visible':warnings.append('Hidden sheet retained: '+sheet.title)
        finally:book.close()
    elif fmt=='docx':
        from docx import Document
        with archive(data):pass
        document=Document(io.BytesIO(data))
        warnings.append('DOCX body only; headers, footers, notes and text boxes require reviewed export')
        # python-docx doesn't expose deleted/revision text as accepted prose.
        xml=document._element.xml
        if '<w:ins ' in xml or '<w:del ' in xml:raise ValueError('Tracked changes need accepted/reviewed export')
        blocks=[dict(text=p.text,location=dict(paragraph=i)) for i,p in enumerate(document.paragraphs,1) if p.text.strip()]
        for i,t in enumerate(document.tables,1):table('table-'+str(i),[[c.text for c in r.cells] for r in t.rows])
    elif fmt=='pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data),strict=True)
        if reader.is_encrypted or len(reader.pages)>MAX_PAGES:raise ValueError('PDF unsupported/limit')
        root=reader.trailer['/Root'];names=root.get('/Names',{})
        if root.get('/OpenAction') or root.get('/AA') or names.get('/JavaScript') or names.get('/EmbeddedFiles'):raise ValueError('Active/embedded PDF unsupported')
        for i,p in enumerate(reader.pages,1):
            text=p.extract_text(extraction_mode='plain')
            if not text or not text.strip():raise ValueError('Scanned/unreadable PDF; reliable OCR unavailable')
            blocks.append(dict(text=text,location=dict(page=i)))
        warnings.append('PDF reading order and tables require review')
    elif fmt=='text':
        blocks=[dict(text=data.decode('utf-8'),location=dict(block=1))]
    else:raise ValueError('Unsupported document format')
    if sum(len(b['text']) for b in blocks)>MAX_BYTES or not tables and not blocks:raise ValueError('Unreadable/oversized source')
    count=sum(len(t['rows']) for t in tables)
    if 'row_count' in meta and (type(meta['row_count']) is not int or meta['row_count']!=count):raise ValueError('Declared population differs')
    if tables and (meta.get('complete_population') is not True or 'row_count' not in meta):warnings.append('Population completeness unverified')
    normalized={'sheets':tables} if tables else {'blocks':[dict(id='block-'+str(i),text=b['text'],**({'page':b['location']['page']} if 'page' in b['location'] else {'section':'paragraph-'+str(b['location'].get('paragraph',i))})) for i,b in enumerate(blocks,1)]}
    source_id=doc['id']+':'+doc['version']
    # Source location remains explicit alongside native normalized field references.
    raw=RawSource(source_id,doc['id'],'normalized_workbook' if tables else 'normalized_document',normalized,meta,raw_ref='sha256:'+hashlib.sha256(data).hexdigest())
    inv=Inventory([raw]);extraction=inv.extractions[source_id].record()
    for f in extraction['fields'].values():
        loc=f['location']
        if tables:loc['row']+=1;loc['sheet']=loc.get('table')
        elif 'block' in loc and loc['block']<=len(blocks):loc.update(blocks[loc['block']-1]['location'])
    result=dict(contract='cao-document/1',id=doc['id'],version=doc['version'],role=doc['role'],
        original_sha256=hashlib.sha256(data).hexdigest(),original_bytes=len(data),original_content_base64=doc['content_base64'],metadata=meta,
        raw_source=dict(id=raw.id,name=raw.name,format=raw.format,payload=raw.payload,metadata=raw.metadata,raw_ref=raw.raw_ref),
        extraction=extraction,tables=tables,blocks=blocks,warnings=warnings,
        qualification='OBSERVATION_ONLY',accounting_authority=False,supersedes=doc.get('supersedes'))
    result['extraction_sha256']=fingerprint(result)
    return result


def ingest(doc, company, scope):
    # Normalize parser-specific malformed-file errors before crossing public API.
    from pypdf.errors import PdfReadError
    from lxml.etree import XMLSyntaxError
    from docx.opc.exceptions import PackageNotFoundError
    try:return _ingest(doc,company,scope)
    except (zipfile.BadZipFile,PdfReadError,XMLSyntaxError,PackageNotFoundError,UnicodeError,EOFError) as exc:
        raise ValueError('Unreadable document format') from None
